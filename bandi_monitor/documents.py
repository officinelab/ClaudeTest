"""Estrazione del testo dagli allegati (PDF, Word, ODT, RTF, ZIP...)."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass

from bs4 import BeautifulSoup


@dataclass
class Documento:
    nome: str
    url: str
    testo: str
    pdf_bytes: bytes | None = None  # PDF senza testo estraibile (scansione): lo legge direttamente Claude
    errore: str | None = None


def _kind(data: bytes, url: str, content_type: str) -> str:
    ct = (content_type or "").lower()
    low = url.lower().split("?")[0]
    if data[:5] == b"%PDF-" or "pdf" in ct:
        return "pdf"
    if data[:2] == b"PK":
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            return "unknown"
        if "word/document.xml" in names:
            return "docx"
        if "content.xml" in names:
            return "odf"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        return "zip"
    if data[:5] == b"{\\rtf":
        return "rtf"
    if "html" in ct or data.lstrip()[:15].lower().startswith((b"<!doctype html", b"<html")):
        return "html"
    if low.endswith(".txt") or ct.startswith("text/"):
        return "text"
    if data[:4] == b"\xd0\xcf\x11\xe0":
        return "ole"  # vecchio .doc/.xls
    return "unknown"


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            parts.append(f"[pag. {i}]\n{page.extract_text() or ''}")
        except Exception:  # pagine corrotte: si prosegue con le altre
            continue
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xml_text(xml: bytes) -> str:
    text = re.sub(rb"<(text:p|text:h|w:p|a:p)[ >/]", b"\n<", xml)
    text = re.sub(rb"<[^>]+>", b" ", text)
    return re.sub(r"[ \t]+", " ", text.decode("utf-8", "replace"))


def _rtf_text(data: bytes) -> str:
    s = data.decode("latin-1", "replace")
    s = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", "replace"), s)
    s = re.sub(r"\\par[d]?", "\n", s)
    s = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", s)
    return s


def _ole_text(data: bytes) -> str:
    # Estrazione grezza delle stringhe di un vecchio .doc: imperfetta ma spesso sufficiente.
    chunks = re.findall(rb"(?:[\x20-\x7e\xa0-\xff]\x00){6,}", data)
    return "\n".join(c.decode("utf-16-le", "ignore") for c in chunks)


def extract_text(data: bytes, url: str, content_type: str = "", nome: str = "", depth: int = 0) -> Documento:
    nome = nome or url.rsplit("/", 1)[-1]
    kind = _kind(data, url, content_type)
    try:
        if kind == "pdf":
            testo = _pdf_text(data)
            # Poco testo rispetto alle pagine: probabilmente una scansione.
            if len(re.sub(r"\[pag\. \d+\]|\s", "", testo)) < 200:
                return Documento(nome, url, testo, pdf_bytes=data)
        elif kind == "docx":
            testo = _docx_text(data)
        elif kind in ("odf", "xlsx"):
            z = zipfile.ZipFile(io.BytesIO(data))
            members = ["content.xml"] if kind == "odf" else [n for n in z.namelist() if n.startswith("xl/sharedStrings")]
            testo = "\n".join(_xml_text(z.read(m)) for m in members)
        elif kind == "zip":
            if depth > 0:
                return Documento(nome, url, "", errore="ZIP annidato ignorato")
            z = zipfile.ZipFile(io.BytesIO(data))
            parts = []
            for info in z.infolist()[:20]:
                if info.is_dir() or info.file_size > 30_000_000:
                    continue
                sub = extract_text(z.read(info), info.filename, nome=info.filename, depth=depth + 1)
                if sub.testo.strip():
                    parts.append(f"--- {info.filename} ---\n{sub.testo}")
            testo = "\n\n".join(parts)
        elif kind == "rtf":
            testo = _rtf_text(data)
        elif kind == "html":
            soup = BeautifulSoup(data, "html.parser")
            for t in soup(["script", "style", "nav", "header", "footer"]):
                t.decompose()
            testo = soup.get_text("\n")
        elif kind == "text":
            testo = data.decode("utf-8", "replace")
        elif kind == "ole":
            testo = _ole_text(data)
        else:
            return Documento(nome, url, "", errore=f"formato non supportato ({content_type or 'sconosciuto'})")
    except Exception as exc:  # un allegato illeggibile non deve fermare l'analisi del bando
        return Documento(nome, url, "", errore=f"{type(exc).__name__}: {exc}")
    testo = re.sub(r"\n{3,}", "\n\n", testo).strip()
    return Documento(nome, url, testo)
